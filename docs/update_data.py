#!/usr/bin/env python3
"""Regenerate the DATA object inside index.html from the Excel recap workbook.

Every data section shown on the page is derived from the workbook:

    1. REKAP NILAI           -> rekap, meta (rata-rata)
    2. NILAI KOMPONEN        -> per_komponen, komponen (bobot)
    3. REKAP KETIDAKLENGKAPAN-> ketidaklengkapan
    4. CEK MANUAL GURU       -> cek_manual
    5. DATA REKAPAN          -> data_rekapan
    7. NILAI TAMBAHAN RESOL. -> bonus, bonus_rules, meta.gambar
    6. METODE & RUBRIK       -> metode, meta.revisi
    Detail - <nama>          -> detail (bukti per komponen)

Formatting is preserved: `rekap` stays pretty-printed, every other section
stays on a single compact line, and the surrounding HTML is untouched.

Usage:
    python3 docs/update_data.py          # write index.html
    python3 docs/update_data.py --check  # report differences, write nothing
"""

import json
import pathlib
import re
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / 'REKAP NILAIAN ASTS - Personal Branding XI DKV 2026-2027.xlsx'
HTML_PATH = ROOT / 'index.html'

SH_REKAP = '1. REKAP NILAI'
SH_KOMPONEN = '2. NILAI KOMPONEN'
SH_KETIDAK = '3. REKAP KETIDAKLENGKAPAN'
SH_CEK = '4. CEK MANUAL GURU'
SH_REKAPAN = '5. DATA REKAPAN'
SH_METODE = '6. METODE & RUBRIK'
SH_BONUS = '7. NILAI TAMBAHAN RESOLUSI'
DETAIL_PREFIX = 'Detail - '

# Sections kept multi-line inside the DATA literal (matches existing file).
PRETTY_SECTIONS = {'rekap'}


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def norm(value):
    """Same normalisation the page uses to match student names."""
    return re.sub(r'[^A-Z0-9]', '', str(value or '').upper())


def txt(value, default=''):
    """Cell -> string, preserving integers as '5' and floats as '2.5'."""
    if value is None:
        return default
    if isinstance(value, bool):
        return default
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def num(value):
    """Parse Excel's '+1,0' / '-' bonus notation into a float."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    raw = str(value).strip().replace('+', '').replace(',', '.')
    if raw in ('', '-', '--'):
        return 0.0
    try:
        return float(raw)
    except ValueError:
        return 0.0


def rows_of(ws, header_row, max_col=None):
    """Yield non-empty data rows below `header_row`."""
    max_col = max_col or ws.max_column
    for row in ws.iter_rows(min_row=header_row + 1,
                            max_col=max_col, values_only=True):
        if all(cell is None for cell in row):
            continue
        if row[0] is None:
            continue
        yield row


def find_header_row(ws, *labels):
    """Row whose cells include every label (exact, case-insensitive)."""
    wanted = [lab.lower() for lab in labels]
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        present = {str(c).strip().lower() for c in row if c is not None}
        if all(lab in present for lab in wanted):
            return i
    raise SystemExit('Header %r not found in sheet %s' % (labels, ws.title))


def header_index(ws, header_row):
    """Map header label -> 0-based column index."""
    return {str(c).strip(): i
            for i, c in enumerate(next(ws.iter_rows(min_row=header_row,
                                                     max_row=header_row,
                                                     values_only=True)))
            if c is not None}


# --------------------------------------------------------------------------- #
# section builders
# --------------------------------------------------------------------------- #
def build_rekap(wb):
    """Sheet 1 -> rekap rows. `bonus` is numeric so the page can compare it."""
    ws = wb[SH_REKAP]
    header_row = find_header_row(ws, 'Nama Siswa')
    col = header_index(ws, header_row)
    out = []
    for row in rows_of(ws, header_row):
        out.append({
            'no': row[col['No']],
            'nama': txt(row[col['Nama Siswa']]),
            'kelas': txt(row[col['Kelas']]),
            'kirim': txt(row[col['Waktu Kirim']]),
            'url': txt(row[col['Link Blogger']]),
            'mockup': txt(row[col['Mockup']]),
            'shotlist': txt(row[col['Shotlist']]),
            'storyboard': txt(row[col['Storyboard']]),
            'mascot': txt(row[col['Mascot']]),
            'prompt': txt(row[col['Prompt']]),
            'kata': txt(row[col['Jml Kata Deskripsi']]),
            'resolusi': txt(row[col['Resolusi (px)']], 'tidak diukur'),
            'rubrik': row[col['Nilai Rubrik /100']],
            'bonus': num(row[col['Nilai Tambah Resolusi']]),
            'nilai': row[col['NILAI AKHIR /100']],
            'kategori': txt(row[col['Kategori']]),
        })
    return out


def build_per_komponen(wb, keys):
    """Sheet 2 -> per-komponen scores (skor 0-4 per komponen)."""
    ws = wb[SH_KOMPONEN]
    header_row = find_header_row(ws, 'Nama', 'NILAI AKHIR', 'Kategori')
    col = header_index(ws, header_row)
    out = []
    for row in rows_of(ws, header_row):
        if not isinstance(row[1], str) or 'Bobot' in row[1]:
            continue
        rec = {
            'no': row[col['No']],
            'nama': txt(row[col['Nama']]),
            'kelas': txt(row[col['Kelas']]),
        }
        for key, idx in keys.items():
            rec[key] = row[idx]
        rec['rubrik'] = row[col['Nilai Rubrik']]
        rec['bonus'] = num(row[col['Nilai Tambah Resolusi']])
        akhir = row[col['NILAI AKHIR']]
        rec['akhir'] = akhir
        rec['nilai'] = akhir
        rec['kategori'] = txt(row[col['Kategori']])
        out.append(rec)
    return out


def build_ketidaklengkapan(wb):
    ws = wb[SH_KETIDAK]
    header_row = find_header_row(ws, 'Nama', 'Komponen', 'Ketentuan Soal')
    col = header_index(ws, header_row)
    return [{
        'nama': txt(row[col['Nama']]),
        'kelas': txt(row[col['Kelas']]),
        'komponen': txt(row[col['Komponen']]),
        'ketentuan': txt(row[col['Ketentuan Soal']]),
        'hasil': txt(row[col['Hasil Aktual']]),
        'catatan': txt(row[col['Kekurangan / Catatan']]),
    } for row in rows_of(ws, header_row)]


def build_cek_manual(wb):
    ws = wb[SH_CEK]
    header_row = find_header_row(ws, 'Nama', 'Butir yang perlu dicek', 'Petunjuk OCR / bukti')
    col = header_index(ws, header_row)
    return [{
        'no': row[col['No']],
        'nama': txt(row[col['Nama']]),
        'butir': txt(row[col['Butir yang perlu dicek']]),
        'bukti': txt(row[col['Petunjuk OCR / bukti']]),
    } for row in rows_of(ws, header_row)]


def build_data_rekapan(wb):
    ws = wb[SH_REKAPAN]
    header_row = find_header_row(ws, 'No', 'Timestamp', 'Link Postingan Blogger')
    col = header_index(ws, header_row)
    return [{
        'no': row[col['No']],
        'ts': txt(row[col['Timestamp']]),
        'kelas': txt(row[col['Kelas']]),
        'nama': txt(row[col['Nama']]),
        'url': txt(row[col['Link Postingan Blogger']]),
        'status': txt(row[col['Status Blogger']]),
        'catatan': txt(row[col['Catatan']]),
    } for row in rows_of(ws, header_row)]


def build_bonus(wb):
    """Sheet 7 -> bonus table + the printed rules block underneath it."""
    ws = wb[SH_BONUS]
    header_row = find_header_row(ws, 'Nama', 'Jumlah Gambar', 'Lebar Maks (px)')
    records = []
    rules = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        first = row[0]
        if first is None:
            continue
        if isinstance(first, (int, float)):
            records.append({
                'no': first,
                'nama': txt(row[1]),
                'jml': txt(row[2]),
                'ge1000': txt(row[3]),
                'maks': txt(row[4]),
                'bonus': txt(row[5], '-'),
                'ket': txt(row[6]),
            })
        elif isinstance(first, str) and first.strip():
            rules.append(first.strip())
    return records, rules


def build_metode(wb, kiriman, gambar):
    """Sheet 6 -> label/value pairs, blank rows dropped (layout uses gaps).

    The sheet still contains the counts from an earlier revision, so the
    factual totals are refreshed to the current number of students/images.
    Historical wording (e.g. '18 siswa dihitung ulang' under Revisi-3) is
    left alone.
    """
    ws = wb[SH_METODE]
    out = []
    for row in ws.iter_rows(max_col=2, values_only=True):
        cells = [c for c in row if c is not None and str(c).strip()]
        if not cells:
            continue
        text = [str(c).strip() for c in cells]
        for i, cell in enumerate(text):
            cell = cell.replace('18 kiriman', '%d kiriman' % kiriman)
            cell = cell.replace('18 dari 18', '%d dari %d' % (kiriman, kiriman))
            cell = cell.replace('119 berkas gambar', '%d berkas gambar' % gambar)
            text[i] = cell
        out.append(text)
    return out


def build_detail(wb, names):
    """One entry per `Detail - <nama>` sheet, keyed by the rekap name.

    Excel truncates sheet names to 31 chars, so the sheet suffix cannot be
    used as a key: the `norm()` lookup in the page would fail for names such
    as 'QUINSYA RAHMANESA SOLEHA'. Each sheet is resolved back to the full
    name from sheet 1 and keyed on that instead.
    """
    out = {}
    for title in wb.sheetnames:
        if not title.startswith(DETAIL_PREFIX):
            continue
        suffix = title[len(DETAIL_PREFIX):]
        candidates = [n for n in names if norm(n).startswith(norm(suffix))]
        if len(candidates) != 1:
            print('  ! sheet %r tidak bisa dipetakan ke nama siswa, dilewati'
                  % title, file=sys.stderr)
            continue
        key = candidates[0]

        ws = wb[title]
        link = txt(ws.cell(2, 1).value)
        url = link.split('Link Blogger:', 1)[-1].strip() \
            if 'Link Blogger:' in link else ''

        # Rows 1-5 are the title, link, time, blank and table header; the
        # actual assessment rows start after the 'Komponen' header.
        header_row = find_header_row(ws, 'Komponen', 'Status',
                                     'Bukti / Hasil Pemeriksaan')
        items = []
        footer = {}
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            label = row[1] if len(row) > 1 else None
            if isinstance(label, str) and label.strip() in ('NILAI RUBRIK',
                                                            'NILAI AKHIR'):
                footer[label.strip()] = txt(row[5])
                continue
            first = row[0]
            if not isinstance(first, (int, float)) or isinstance(first, bool):
                if first != '+':
                    continue
            items.append({
                'no': first,
                'komponen': txt(row[1]),
                'ketentuan': txt(row[2]),
                'status': txt(row[3], '-'),
                'bukti': txt(row[4]),
                'skor': txt(row[5]),
                'bobot': txt(row[6]),
                'nilai': txt(row[7]),
            })

        out[key] = {
            'nama': key,
            'items': items,
            'meta': {'url': url, 'waktu': txt(ws.cell(3, 1).value)},
            'footer': footer,
        }
    return out


def build_meta(wb, rekap, gambar):
    """Refresh the derived summary numbers; keep the hand-written wording."""
    ws = wb[SH_METODE]
    revisi, revisi_tanggal = '', ''
    for row in ws.iter_rows(max_col=2, values_only=True):
        label = row[0]
        if isinstance(label, str) and re.match(r'^REVISI RUBRIK KE-\d', label):
            revisi = label.split('(')[0].strip()
            tail = label.split('(', 1)[1] if '(' in label else ''
            revisi_tanggal = tail.rstrip(') ').strip()
            break

    nilai = [float(r['nilai']) for r in rekap]
    rubrik = [float(r['rubrik']) for r in rekap]
    bonus = [float(r['bonus']) for r in rekap]
    n = len(rekap)

    return {
        'judul': 'Rekap Penilaian Projek ASTS - Personal Branding Creative Campaign',
        'mapel': 'DKV - AI dalam Desain',
        'semester': 'Semester 1',
        'tp': 'TP 2026/2027',
        'sekolah': 'SMK Negeri 9 Garut',
        'kelas': 'XI DKV 1-4',
        'batas': 'Jumat, 2 Oktober 2026 pukul 23.59 WIB',
        'sumber': ('rekapan Google Form (%d kiriman) + akses langsung ke Blogger'
                   ' + %d berkas gambar (OCR & dimensi piksel asli)'
                   % (n, gambar)),
        'gambar': gambar,
        'revisi': revisi,
        'revisiTanggal': revisi_tanggal,
        'rata': round(sum(nilai) / n, 2),
        'rataRubrik': round(sum(rubrik) / n, 2),
        'rataBonus': round(sum(bonus) / n, 2),
    }


# --------------------------------------------------------------------------- #
# DATA serialisation
# --------------------------------------------------------------------------- #
def data_span(text):
    """Return (start, end) of the object literal assigned to `const DATA`."""
    marker = 'const DATA = '
    i = text.index(marker) + len(marker)
    depth = 0
    in_str = False
    escaped = False
    for k in range(i, len(text)):
        c = text[k]
        if in_str:
            if escaped:
                escaped = False
            elif c == '\\':
                escaped = True
            elif c == '"':
                in_str = False
            continue
        if c == '"':
            in_str = True
        elif c == '{':
            depth += 1
        elif c == '}':
            depth -= 1
            if depth == 0:
                return i, k + 1
    raise SystemExit('Could not locate the end of the DATA object')


def serialise(data):
    """Render DATA back to text using the file's existing layout."""
    parts = []
    for key, value in data.items():
        if key in PRETTY_SECTIONS:
            body = json.dumps(value, ensure_ascii=False, indent=2)
        else:
            body = json.dumps(value, ensure_ascii=False,
                              separators=(', ', ': '))
        parts.append('%s: %s' % (json.dumps(key, ensure_ascii=False), body))
    return '{' + ', '.join(parts) + '}'


# --------------------------------------------------------------------------- #
def main():
    if not EXCEL_PATH.exists():
        raise SystemExit('Workbook not found: %s' % EXCEL_PATH)
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    ws_komponen = wb[SH_KOMPONEN]
    head = next(ws_komponen.iter_rows(min_row=4, max_row=4, values_only=True))
    keys, bobot = {}, {}
    for i, label in enumerate(head):
        if not isinstance(label, str):
            continue
        m = re.match(r'^(.+?)\s*\(bobot\s+(\d+)\)$', label.strip())
        if m:
            name, weight = m.group(1).strip(), int(m.group(2))
            key = 'k%d' % (len(bobot) + 1)
            keys[key] = i
            bobot[key] = {'key': key, 'col': len(bobot) + 3,
                          'nama': name, 'bobot': weight}
    if len(bobot) != 12:
        raise SystemExit('Expected 12 komponen, found %d' % len(bobot))

    rekap = build_rekap(wb)
    bonus, bonus_rules = build_bonus(wb)
    gambar = sum(int(b['jml']) for b in bonus if b['jml'].isdigit())
    metode = build_metode(wb, len(rekap), gambar)
    per_komponen = build_per_komponen(wb, keys)
    detail = build_detail(wb, [r['nama'] for r in rekap])

    text = HTML_PATH.read_text(encoding='utf-8')
    start, end = data_span(text)
    old = json.loads(text[start:end])

    new = dict(old)
    new['meta'] = build_meta(wb, rekap, gambar)
    new['komponen'] = list(bobot.values())
    new['rekap'] = rekap
    new['per_komponen'] = per_komponen
    new['ketidaklengkapan'] = build_ketidaklengkapan(wb)
    new['cek_manual'] = build_cek_manual(wb)
    new['data_rekapan'] = build_data_rekapan(wb)
    new['bonus'] = bonus
    new['bonus_rules'] = bonus_rules
    new['metode'] = metode
    new['detail'] = detail

    if '--check' in sys.argv:
        print('%-18s %8s %8s' % ('section', 'before', 'after'))
        for key in new:
            b = json.dumps(old.get(key), sort_keys=True)
            a = json.dumps(new.get(key), sort_keys=True)
            if key == 'meta':
                for mk in ('sumber', 'gambar', 'revisi', 'revisiTanggal',
                           'rata', 'rataRubrik', 'rataBonus'):
                    if old['meta'].get(mk) != new['meta'].get(mk):
                        print('  meta.%-14s %r -> %r'
                              % (mk, old['meta'].get(mk), new['meta'].get(mk)))
            size = lambda v: len(v) if isinstance(v, (list, dict)) else 1
            print('%-18s %8d %8d %s'
                  % (key, size(old.get(key)), size(new.get(key)),
                     'OK' if a == b else 'BERUBAH'))
        return 0

    out = text[:start] + serialise(new) + text[end:]
    HTML_PATH.write_text(out, encoding='utf-8')
    print('index.html diperbarui dari Excel:')
    print('  rekap            %d siswa' % len(rekap))
    print('  per_komponen     %d siswa' % len(per_komponen))
    print('  detail           %d siswa' % len(new['detail']))
    print('  data_rekapan     %d baris' % len(new['data_rekapan']))
    print('  bonus            %d baris' % len(bonus))
    print('  ketidaklengkapan %d temuan' % len(new['ketidaklengkapan']))
    print('  cek_manual       %d butir' % len(new['cek_manual']))
    print('  rata-rata akhir  %s (rubrik %s + nilai tambah %s)'
          % (new['meta']['rata'], new['meta']['rataRubrik'],
             new['meta']['rataBonus']))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())